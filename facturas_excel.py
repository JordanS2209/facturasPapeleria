from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook

IVA_RATE = 0.15
DEFAULT_OUTPUT = "factura_con_iva.xlsx"
SUMMARY_LABELS = {"Subtotal", "IVA total", "Total general"}
HEADERS = [
    "Cantidad",
    "Descripción del producto",
    "Precio unitario sin IVA",
    "Precio unitario con IVA",
    "Valor total sin IVA",
    "IVA del 15%",
    "Total con IVA",
    "Observaciones",
]
COMPARISON_HEADERS = [
    "Producto solicitado",
    "Cantidad solicitada",
    "Cantidad recibida",
    "Diferencia",
    "Estado",
    "Observaciones",
]


@dataclass
class InvoiceItem:
    cantidad: float | None
    descripcion: str | None
    precio_base: float | None
    observaciones: list[str]


def round2(value: float) -> float:
    return round(value + 1e-9, 2)


def normalize_number(token: str) -> float:
    cleaned = token.replace("$", "").replace(" ", "")
    if cleaned.count(",") and cleaned.count("."):
        cleaned = cleaned.replace(".", "").replace(",", ".")
    else:
        cleaned = cleaned.replace(",", ".")
    return float(cleaned)


def _is_image(path: Path) -> bool:
    return path.suffix.lower() in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}


def extract_text(path: Path) -> str:
    if _is_image(path):
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError(
                "Faltan dependencias OCR. Instala pillow y pytesseract para procesar fotos."
            ) from exc

        return pytesseract.image_to_string(Image.open(path), lang="spa+eng")

    return path.read_text(encoding="utf-8")


def parse_invoice_text(text: str) -> tuple[list[InvoiceItem], list[str]]:
    items: list[InvoiceItem] = []
    unreadable: list[str] = []

    for line in [ln.strip() for ln in text.splitlines() if ln.strip()]:
        lowered = line.lower()
        if any(tag in lowered for tag in ("subtotal", "total", "iva")):
            continue

        cantidad = None
        descripcion = None
        precio_base = None
        observaciones: list[str] = []

        qty_match = re.match(r"^(\d+(?:[\.,]\d+)?)\b", line)
        if qty_match:
            try:
                cantidad = normalize_number(qty_match.group(1))
            except ValueError:
                pass

        number_tokens = re.findall(r"\d+(?:[\.,]\d+)?", line)
        numeric_values: list[float] = []
        for token in number_tokens:
            try:
                numeric_values.append(normalize_number(token))
            except ValueError:
                continue

        if len(numeric_values) >= 2:
            precio_base = numeric_values[-2]
        elif len(numeric_values) == 1:
            if cantidad and cantidad > 0:
                precio_base = numeric_values[0] / cantidad
            else:
                precio_base = numeric_values[0]

        desc_candidate = line
        if qty_match:
            desc_candidate = desc_candidate[qty_match.end() :].strip(" -:")
        for token in number_tokens:
            desc_candidate = re.sub(rf"\b{re.escape(token)}\b", "", desc_candidate)
        desc_candidate = re.sub(r"\s+", " ", desc_candidate).strip(" -:")
        if desc_candidate:
            descripcion = desc_candidate

        if cantidad is None:
            observaciones.append("Cantidad no legible")
            unreadable.append(f"No se pudo leer cantidad en: '{line}'")
        if descripcion is None:
            observaciones.append("Descripción no legible")
            unreadable.append(f"No se pudo leer descripción en: '{line}'")
        if precio_base is None:
            observaciones.append("Precio unitario no legible")
            unreadable.append(f"No se pudo leer precio unitario en: '{line}'")

        items.append(
            InvoiceItem(
                cantidad=cantidad,
                descripcion=descripcion,
                precio_base=precio_base,
                observaciones=observaciones,
            )
        )

    return items, unreadable


def build_invoice_row(item: InvoiceItem, prices_include_iva: bool = False) -> list[float | str]:
    cantidad = item.cantidad
    descripcion = item.descripcion or "REVISAR"
    observacion = "; ".join(item.observaciones)

    if cantidad is None or item.precio_base is None:
        return [
            cantidad if cantidad is not None else "REVISAR",
            descripcion,
            "REVISAR",
            "REVISAR",
            "REVISAR",
            "REVISAR",
            "REVISAR",
            observacion or "REVISAR",
        ]

    if prices_include_iva:
        precio_unitario_con_iva = round2(item.precio_base)
        precio_unitario_sin_iva = round2(precio_unitario_con_iva / (1 + IVA_RATE))
    else:
        precio_unitario_sin_iva = round2(item.precio_base)
        precio_unitario_con_iva = round2(precio_unitario_sin_iva * (1 + IVA_RATE))

    valor_total_sin_iva = round2(cantidad * precio_unitario_sin_iva)
    iva_15 = round2(valor_total_sin_iva * IVA_RATE)
    total_con_iva = round2(valor_total_sin_iva + iva_15)

    return [
        round2(cantidad),
        descripcion,
        precio_unitario_sin_iva,
        precio_unitario_con_iva,
        valor_total_sin_iva,
        iva_15,
        total_con_iva,
        observacion,
    ]


def _is_summary_row_label(value: object) -> bool:
    return isinstance(value, str) and value.strip() in SUMMARY_LABELS


def _to_float(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    return None


def prepare_invoice_sheet(workbook: Workbook):
    sheet = workbook.active
    sheet.title = "Facturas"
    if sheet.max_row == 1 and all(sheet.cell(1, c).value is None for c in range(1, len(HEADERS) + 1)):
        for index, header in enumerate(HEADERS, start=1):
            sheet.cell(1, index, header)

    # Limpia resumen anterior si existe
    row = sheet.max_row
    while row > 1:
        label = sheet.cell(row, 2).value
        if _is_summary_row_label(label):
            sheet.delete_rows(row, 1)
        row -= 1

    return sheet


def append_invoice_rows(sheet, rows: Iterable[list[float | str]]):
    for row in rows:
        sheet.append(row)


def add_summary_rows(sheet):
    subtotal = 0.0
    iva_total = 0.0
    total_general = 0.0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_summary_row_label(row[1]):
            continue
        total_sin_iva = _to_float(row[4])
        iva = _to_float(row[5])
        total = _to_float(row[6])
        if total_sin_iva is not None:
            subtotal += total_sin_iva
        if iva is not None:
            iva_total += iva
        if total is not None:
            total_general += total

    sheet.append([None, "Subtotal", None, None, round2(subtotal), None, None, None])
    sheet.append([None, "IVA total", None, None, None, round2(iva_total), None, None])
    sheet.append([None, "Total general", None, None, None, None, round2(total_general), None])


def normalize_product_name(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower())


def parse_order_text(text: str) -> list[tuple[str, float]]:
    result: list[tuple[str, float]] = []
    for line in [ln.strip() for ln in text.splitlines() if ln.strip()]:
        match = re.match(r"^(\d+(?:[\.,]\d+)?)\s+(.+)$", line)
        if not match:
            continue
        qty = normalize_number(match.group(1))
        product = match.group(2).strip(" -:")
        if product:
            result.append((product, qty))
    return result


def parse_order_file(path: Path) -> list[tuple[str, float]]:
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        parsed: list[tuple[str, float]] = []
        for row in data if isinstance(data, list) else []:
            product = str(row.get("producto", "")).strip()
            cantidad = row.get("cantidad")
            if product and cantidad is not None:
                parsed.append((product, float(cantidad)))
        return parsed

    if path.suffix.lower() == ".csv":
        parsed = []
        with path.open("r", encoding="utf-8", newline="") as file:
            reader = csv.DictReader(file)
            for row in reader:
                product = (
                    row.get("producto")
                    or row.get("Producto")
                    or row.get("descripcion")
                    or row.get("Descripción")
                    or ""
                ).strip()
                qty_raw = row.get("cantidad") or row.get("Cantidad")
                if not product or not qty_raw:
                    continue
                parsed.append((product, normalize_number(qty_raw)))
        return parsed

    text = extract_text(path)
    return parse_order_text(text)


def build_received_map(sheet) -> dict[str, float]:
    received: dict[str, float] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if _is_summary_row_label(row[1]):
            continue
        qty = _to_float(row[0])
        name = row[1] if isinstance(row[1], str) else None
        if qty is None or not name or name == "REVISAR":
            continue
        key = normalize_product_name(name)
        received[key] = round2(received.get(key, 0.0) + qty)
    return received


def update_comparison_sheet(workbook: Workbook, requested: list[tuple[str, float]], received: dict[str, float]):
    if "Comparación" in workbook.sheetnames:
        del workbook["Comparación"]

    sheet = workbook.create_sheet("Comparación")
    sheet.append(COMPARISON_HEADERS)

    for product, qty_requested in requested:
        key = normalize_product_name(product)
        qty_received = received.get(key, 0.0)
        difference = round2(qty_requested - qty_received)

        if qty_received == 0:
            status = "Faltante"
        elif qty_received >= qty_requested:
            status = "Completo"
        else:
            status = "Parcial"

        obs = ""
        if qty_received == 0:
            obs = "Producto no identificado en la factura"
        elif qty_received > qty_requested:
            obs = "Se recibió más de lo solicitado"

        sheet.append([
            product,
            round2(qty_requested),
            round2(qty_received),
            difference,
            status,
            obs,
        ])


def run(invoice_paths: list[Path], output: Path, order_path: Path | None, prices_include_iva: bool):
    workbook = load_workbook(output) if output.exists() else Workbook()
    sheet = prepare_invoice_sheet(workbook)

    parsed_rows: list[list[float | str]] = []
    unreadable: list[str] = []

    for invoice_path in invoice_paths:
        text = extract_text(invoice_path)
        items, issues = parse_invoice_text(text)
        unreadable.extend(issues)
        parsed_rows.extend(build_invoice_row(item, prices_include_iva=prices_include_iva) for item in items)

    append_invoice_rows(sheet, parsed_rows)
    add_summary_rows(sheet)

    if order_path:
        requested = parse_order_file(order_path)
        received_map = build_received_map(sheet)
        update_comparison_sheet(workbook, requested=requested, received=received_map)

    workbook.save(output)
    return unreadable


def main():
    parser = argparse.ArgumentParser(description="Procesa facturas y genera un Excel con IVA.")
    parser.add_argument("facturas", nargs="+", help="Rutas de imágenes o texto de facturas")
    parser.add_argument("--pedido", help="Ruta opcional de lista/foto del pedido")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Archivo Excel destino")
    parser.add_argument(
        "--precios-con-iva",
        action="store_true",
        help="Interpretar precios detectados como valores con IVA incluido",
    )

    args = parser.parse_args()
    invoice_paths = [Path(path).resolve() for path in args.facturas]
    output = Path(args.output).resolve()
    order_path = Path(args.pedido).resolve() if args.pedido else None

    unreadable = run(
        invoice_paths=invoice_paths,
        output=output,
        order_path=order_path,
        prices_include_iva=args.precios_con_iva,
    )

    print(f"Archivo actualizado: {output}")
    if unreadable:
        print("Partes no legibles (REVISAR):")
        for item in unreadable:
            print(f"- {item}")


if __name__ == "__main__":
    main()
